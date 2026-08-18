# -*- coding: utf-8 -*-
"""연구실 내부 공유용 R0 (2026-08-12). 제출용(Downloads 템플릿)과 데이터는 동일 소스(NO2=콜드
이상치제거, ANs=g=0.82)지만, 공란 대신 숫자 flag 컬럼(0/1/2/3/4/9)을 달아서 왜 빠졌는지/왜
의심스러운지 남겨둔다. PNs는 콜드+핫 NO2 인젝션 재실험 후 추가 예정, 이번엔 미포함.

flag 근거가 된 유지보수/인젝션 이벤트 타임스탬프는 docs/field_log_summary_2026.md 필드로그
전수조사 결과에서 가져옴 -- 새 캠페인이면 이 창들은 무의미하니 다시 뽑아야 함.

재실행하면 Output/figure/NIER_submit/SDD_CAESAR_O3campaign_2026_labshare.xlsx 를 새로 만든다
(매번 openpyxl.Workbook()으로 새로 생성 -- 기존 파일이 있어도 덮어씀, 수기 수정 보존 안 됨).
"""
import os, sys
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl
import numpy as np
import pandas as pd

SUB = r"C:\Doasis_Work\Output\figure\NIER_submit"
NO2_CSV = os.path.join(SUB, "GIST_CAESAR_cold_NO2_5min_KST_20260517_20260618.csv")
ANS_CSV = os.path.join(SUB, "GIST_CAESAR_ANs_5min_KST_20260518_20260710.csv")

OUT_XLSX = r"C:\Doasis_Work\Output\figure\NIER_submit\SDD_CAESAR_O3campaign_2026_labshare.xlsx"

COLD_BAD_WINDOWS_9 = [
    ("2026-05-20 00:00:00", "2026-05-24 00:00:00"),
    ("2026-05-27 18:00:00", "2026-05-28 12:00:00"),
    ("2026-06-05 00:00:00", "2026-06-05 23:59:59"),
]
COLD_CAL_WINDOWS_4 = [
    ("2026-06-02 16:10:00", "2026-06-02 16:55:00"),
]

HOT_SUSPECT_3 = [
    ("2026-05-22 09:00:00", "2026-05-22 10:22:00"),
    ("2026-05-23 08:49:00", "2026-05-23 09:29:00"),
    ("2026-05-24 10:15:00", "2026-05-24 10:55:00"),
    ("2026-05-26 17:18:00", "2026-05-26 18:15:00"),
    ("2026-06-05 21:17:00", "2026-06-05 22:34:00"),
    ("2026-06-07 16:49:00", "2026-06-07 17:29:00"),
    ("2026-06-09 12:46:00", "2026-06-09 13:00:00"),
    ("2026-06-12 04:21:00", "2026-06-12 05:01:00"),
    ("2026-06-14 11:27:00", "2026-06-14 12:07:00"),
    ("2026-06-21 06:38:00", "2026-06-21 22:38:00"),
    ("2026-06-22 09:08:00", "2026-06-22 09:48:00"),
    ("2026-06-22 10:12:00", "2026-06-22 10:52:00"),
    ("2026-06-23 17:01:00", "2026-06-23 17:41:00"),
    ("2026-07-09 08:40:00", "2026-07-09 09:31:00"),
]
HOT_MISSING_9 = [
    ("2026-06-09 11:00:00", "2026-06-09 12:46:00"),
]
HOT_CAL_4 = [
    ("2026-06-09 19:55:00", "2026-06-09 20:45:00"),
]


def in_any(t, windows):
    for s, e in windows:
        if pd.Timestamp(s) <= t <= pd.Timestamp(e):
            return True
    return False


no2 = pd.read_csv(NO2_CSV, comment="#")
no2["datetime_KST"] = pd.to_datetime(no2["datetime_KST"])
no2["v"] = pd.to_numeric(no2["NO2_ppb"], errors="coerce")

def cold_flag(row):
    t = row["datetime_KST"]
    if pd.isna(row["v"]):
        return 9
    if in_any(t, COLD_CAL_WINDOWS_4):
        return 4
    if in_any(t, COLD_BAD_WINDOWS_9):
        return 9
    if row["QC_flag"] == "low_quality":
        return 9
    if row["QC_flag"] == "unphysical":
        return 9
    return 0

no2["flag"] = no2.apply(cold_flag, axis=1)
no2.loc[no2["flag"].isin([9]), "v"] = np.nan
no2.loc[no2["flag"] == 4, "v"] = np.nan  # calibration/injection window -- not ambient, blank it too

ans = pd.read_csv(ANS_CSV, comment="#")
ans["datetime_KST"] = pd.to_datetime(ans["datetime_KST"])
ans["v"] = pd.to_numeric(ans["ANs_ppb"], errors="coerce")

def hot_flag(row):
    t = row["datetime_KST"]
    if in_any(t, HOT_CAL_4):
        return 4
    if in_any(t, HOT_MISSING_9):
        return 9
    if pd.isna(row["v"]):
        return 9
    if in_any(t, HOT_SUSPECT_3):
        return 3
    return 0

ans["flag"] = ans.apply(hot_flag, axis=1)
ans.loc[ans["flag"].isin([4, 9]), "v"] = np.nan

t_min = min(no2.datetime_KST.min(), ans.datetime_KST.min())
t_max = max(no2.datetime_KST.max(), ans.datetime_KST.max())
grid = pd.date_range(t_min, t_max, freq="5min")

no2_i = no2.set_index("datetime_KST")[["v", "flag"]].reindex(grid)
no2_i["flag"] = no2_i["flag"].fillna(9).astype(int)
ans_i = ans.set_index("datetime_KST")[["v", "flag"]].reindex(grid)
ans_i["flag"] = ans_i["flag"].fillna(9).astype(int)

full = pd.DataFrame(index=grid)
full["SDD_CAESAR_NO2_ppbv"] = no2_i["v"].round(4)
full["SDD_CAESAR_NO2_flag"] = no2_i["flag"]
full["SDD_CAESAR_ANs_ppbv"] = ans_i["v"].round(4)
full["SDD_CAESAR_ANs_flag"] = ans_i["flag"]
full = full.reset_index().rename(columns={"index": "START_TIME"})
full["END_TIME"] = full["START_TIME"] + pd.Timedelta(minutes=5)
full = full[["START_TIME", "END_TIME", "SDD_CAESAR_NO2_ppbv", "SDD_CAESAR_NO2_flag",
             "SDD_CAESAR_ANs_ppbv", "SDD_CAESAR_ANs_flag"]]

for sp in ("NO2", "ANs"):
    vc = full[f"SDD_CAESAR_{sp}_flag"].value_counts().sort_index()
    print(sp, dict(vc))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Data"
wb.create_sheet("Info")
headers = list(full.columns)
for c, h in enumerate(headers, start=1):
    ws.cell(1, c, h)
for i, row in enumerate(full.itertuples(index=False), start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(i, c, None if (isinstance(v, float) and pd.isna(v)) else v)
print(f"Data: {len(full)} rows x {len(headers)} cols")

FLAG_NOTE = ("Flag codes: 0=Good/Valid, 1=Below LOD, 2=Interpolated, 3=Suspect(instrument "
             "unstable/maintenance nearby, value kept), 4=Calibration(injection test, not "
             "ambient), 9=Missing/Bad. Suspect windows derived from field-log maintenance "
             "events (LED TEC adjustments, filter changes, cylinder swaps, filter test "
             "06-21). Injection contamination confirmed: cold 06-02 16:10-16:55 (raw file "
             "N_valid 5->1 drop), hot 06-09 19:55-20:45 (raw noise spike).")

info = wb["Info"]
LABELS = {
    1: "연구책임자 성명(영문명)", 2: "연구책임자 소속(영문)", 3: "연구책임자 이메일",
    4: "연구책임자 연락처", 5: "자료관리자 성명(영문명)", 6: "자료관리자 이메일",
    7: "자료관리자 연락처", 8: "자료버전 *R0부터 시작", 9: "측정 장소(위도, 경도)",
    10: "측정 장비명 및 분석방법 *측정 항목별로 간단히 작성",
    11: "측정 원자료 시간 해상도 *측정 항목별로 간단히 작성",
    12: "최소검출한계(MDL) *측정 항목 또는 분석법별 작성",
    13: "측정불확도 또는 정밀도 *측정 항목 또는 분석법별 작성",
    14: "기타 정보(자료 처리 방법 등 특이사항)",
    15: "변수명 1", 16: "변수명 2",
}
for r, lab in LABELS.items():
    info.cell(r, 1).value = lab
info.cell(1, 2).value = "민경은 (Kyung-Eun Min)"
info.cell(2, 2).value = "광주과학기술원 (Gwangju Institute of Science and Technology, GIST)"
info.cell(3, 2).value = "kemin@gist.ac.kr"
info.cell(4, 2).value = "062-715-3280"
info.cell(5, 2).value = "이경환 (Gyung-Hwan Lee)"
info.cell(6, 2).value = "gh548080@gist.ac.kr"
info.cell(7, 2).value = "062-715-2470"
info.cell(8, 2).value = "R0"
info.cell(9, 2).value = "전남 순천시 해룡면 신대리 2040 (34.92959, 127.54514)"
info.cell(10, 2).value = ("Nitrogen dioxide (NO2): CAESAR-cold, BBCEAS (Broadband Cavity-Enhanced "
                          "Absorption Spectroscopy)\nTotal alkyl nitrates (ANs): CAESAR-hot, "
                          "TD-BBCEAS (Thermal-Dissociation BBCEAS), g=0.82 채널이득보정 적용")
info.cell(11, 2).value = "Nitrogen dioxide (NO2): 1sec\nTotal alkyl nitrates (ANs): 1sec"
info.cell(12, 2).value = "TBD"
info.cell(13, 2).value = "TBD"
info.cell(14, 2).value = (FLAG_NOTE + "\nPNs(과산화질산염)는 콜드/핫 NO2 인젝션 재실험 진행 후 추가 예정 — "
                          "이번 버전엔 미포함.\n현재 농도 보정이 완료되지 않은 R0 버전 자료로, 데이터 "
                          "사용 시 반드시 PI에게 이메일로 사전 문의 필요.")
info.cell(15, 2).value = "SDD_CAESAR_NO2_ppbv"
info.cell(16, 2).value = "SDD_CAESAR_ANs_ppbv"
info.cell(17, 1).value = None
info.cell(17, 2).value = None

wb.save(OUT_XLSX)
print("saved", OUT_XLSX)
